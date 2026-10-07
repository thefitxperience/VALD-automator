"""
Bodydot payment report — the Bodydot counterpart to "Payment - Month YEAR.xlsx".

Three sheets, styled from the Bodydot monthly report (blue header bands, orange
accent chips, Source Sans Pro / Tahoma):

  REPORT                          totals per branch for the selected month, plus
                                  a per-gym roll-up on the right.
  Body Motions - RUH - Al Sahafa  every dispatched program, oldest month first,
  Body Masters - RUH - Al Aarid   with a green separator band between months.

Bodydot runs at a single branch per gym, so there are only two branch sheets.
There is no TEST VALIDITY sheet here — payment is about programs delivered, and
an invalid test never produces one.

Unlike the VALD payment file, nothing is baked into the template: Bodydot only
started in May 2026, so the entire history is rebuilt from `bodydot_tests` on
every generation and the template stays an empty shell.
"""
import calendar
import io
import os
from copy import copy
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from bodydot_report_generator import _client_display_name, _client_ref

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(BASE_DIR, "Bodydot Payment - Month YEAR.xlsx")

REPORT_SHEET = "REPORT"
# Gym → (sheet name, the REPORT cell holding that gym's monthly total)
GYM_SHEETS: dict[str, tuple[str, str]] = {
    "Body Masters": ("Body Masters - RUH - Al Aarid", "B9"),
    "Body Motions": ("Body Motions - RUH - Al Sahafa", "E9"),
}

# Bodydot's first month of data — nothing before this can appear in the file.
START_YEAR_MONTH = (2026, 5)

# ── Month separator band, same convention as the VALD payment file ───────────
GREEN_FILL = PatternFill(fill_type="solid", fgColor="8CC075")
WHITE_FILL = PatternFill(fill_type="solid", fgColor="FFFFFF")
_WHITE_SIDE = Side(style="thin", color="FFFFFF")
GREEN_BORDER = Border(top=_WHITE_SIDE, bottom=_WHITE_SIDE,
                      left=Side(style=None), right=Side(style=None))
SEPARATOR_HEIGHT = 40
DATA_ROW_HEIGHT = 64.5          # matches the Bodydot report's data rows
DATE_FORMAT = "DD/MM/YYYY"
FIRST_DATA_ROW = 7

LATE_UPLOAD_FONT = Font(name="Source Sans Pro", size=20)
LATE_UPLOAD_ALIGN = Alignment(horizontal="left", vertical="center")


def _as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        try:
            return date.fromisoformat(v[:10])
        except ValueError:
            return None
    return None


def _months_through(year: int, month: int) -> list[tuple[int, int]]:
    """Every (year, month) from Bodydot's start through the selected month."""
    out, (y, m) = [], START_YEAR_MONTH
    while (y, m) <= (year, month):
        out.append((y, m))
        m += 1
        if m > 12:
            y, m = y + 1, 1
    return out


def _copy_row_style(ws, src_row: int, dst_row: int, max_col: int):
    for col in range(1, max_col + 1):
        src, dst = ws.cell(src_row, col), ws.cell(dst_row, col)
        if isinstance(src, MergedCell) or isinstance(dst, MergedCell):
            continue
        if src.has_style:
            dst._style = copy(src._style)


def _band(ws, row: int, fill, border=None):
    """Paint one separator row across the five data columns."""
    for col in range(1, 6):
        cell = ws.cell(row, col)
        if isinstance(cell, MergedCell):
            continue
        cell.value = None
        cell.fill = copy(fill)
        if border is not None:
            cell.border = copy(border)
    ws.row_dimensions[row].height = SEPARATOR_HEIGHT


def generate_bodydot_payment_report(
    tests: list[dict],          # approved, valid, non-ignored bodydot_tests rows
    month: int,
    year: int,
    report_date: date | None = None,
) -> bytes:
    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(f"Bodydot payment template not found: {TEMPLATE_PATH}")

    months = _months_through(year, month)
    if not months:
        start = f"{calendar.month_name[START_YEAR_MONTH[1]]} {START_YEAR_MONTH[0]}"
        raise ValueError(
            f"{calendar.month_name[month]} {year} is before Bodydot started ({start})."
        )

    # Bucket every dispatched program by (gym, dispatch year-month).
    by_gym_month: dict[tuple[str, tuple[int, int]], list[dict]] = {}
    for t in tests:
        dd = _as_date(t.get("dispatch_date"))
        if not dd:
            continue
        key = (t.get("gym", ""), (dd.year, dd.month))
        by_gym_month.setdefault(key, []).append(t)
    for rows in by_gym_month.values():
        rows.sort(key=lambda t: (_as_date(t.get("dispatch_date")) or date.min,
                                 _as_date(t.get("test_date")) or date.min))

    wb = load_workbook(TEMPLATE_PATH, data_only=False)
    rpt_date = report_date or date.today()

    for gym, (sheet_name, total_cell) in GYM_SHEETS.items():
        ws = wb[sheet_name]
        ws["B3"] = rpt_date

        # The template ships empty, so clear defensively and start at row 7.
        for row in ws.iter_rows(min_row=FIRST_DATA_ROW, max_row=ws.max_row,
                                max_col=ws.max_column):
            for cell in row:
                if not isinstance(cell, MergedCell):
                    cell.value = None

        row_at = FIRST_DATA_ROW
        wrote_a_month = False
        for ym in months:
            programs = by_gym_month.get((gym, ym), [])
            if not programs:
                continue
            # A separator sits *between* months, never above the first one.
            if wrote_a_month:
                _band(ws, row_at, WHITE_FILL)
                _band(ws, row_at + 1, GREEN_FILL, GREEN_BORDER)
                _band(ws, row_at + 2, WHITE_FILL)
                row_at += 3
            wrote_a_month = True

            for prog in programs:
                _copy_row_style(ws, FIRST_DATA_ROW, row_at, ws.max_column)
                ws.row_dimensions[row_at].height = DATA_ROW_HEIGHT

                test_date = _as_date(prog.get("test_date"))
                dispatch_date = _as_date(prog.get("dispatch_date"))
                ws.cell(row_at, 1, _client_ref(prog))
                ws.cell(row_at, 2, _client_display_name(prog))
                ws.cell(row_at, 3, prog.get("trainer_name") or "")
                ws.cell(row_at, 4, test_date).number_format = DATE_FORMAT
                ws.cell(row_at, 5, dispatch_date).number_format = DATE_FORMAT

                # Col F flags a program dispatched in a later month than its test.
                if test_date and dispatch_date and (
                    (test_date.year, test_date.month) != (dispatch_date.year, dispatch_date.month)
                ):
                    cell = ws.cell(row_at, 6)
                    if not isinstance(cell, MergedCell):
                        cell.value = "Late Upload"
                        cell.font = copy(LATE_UPLOAD_FONT)
                        cell.alignment = copy(LATE_UPLOAD_ALIGN)

                row_at += 1

        wb[REPORT_SHEET][total_cell] = len(by_gym_month.get((gym, (year, month)), []))

    wb[REPORT_SHEET]["B3"] = rpt_date

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
